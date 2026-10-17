from __future__ import annotations

from html import escape
from pathlib import Path
import json
import re
import shutil

from openpyxl import load_workbook


SOURCE = Path(r"X:\My proekt\скоростя\obshaya tablitsa.xlsx")
OUT = SOURCE.parent / "html"

CATEGORIES = {
    "passenger": {
        "title": "Пассажирские",
        "subtitle": "Скорости для пассажирских поездов",
        "columns": [(4, "Главные пути"), (8, "Станция: главные"), (12, "Станция: приёмо-отправочные")],
    },
    "freight": {
        "title": "Грузовые",
        "subtitle": "Скорости для грузовых поездов",
        "columns": [(5, "Главные: грузовые"), (9, "Станция: главный путь"), (13, "Приёмо-отправочные")],
    },
    "empty": {
        "title": "Порожние",
        "subtitle": "Скорости для поездов с порожними вагонами",
        "columns": [(6, "Главные пути"), (10, "Станции")],
    },
    "18-100": {
        "title": "18-100",
        "subtitle": "Порожние вагоны на тележках 18-100",
        "columns": [(7, "Главные пути"), (11, "Станции")],
    },
}

PAGES = {}
for track in ("1", "2"):
    for key, category in CATEGORIES.items():
        filename = f"track-{track}-{key}.html"
        PAGES[filename] = {
            **category,
            "title": f"{track} путь — {category['title']}",
            "kind": "speed_track",
            "track": track,
            "category": key,
        }

SECTIONS = {
    "kor-vih": {"label": "Коршуниха — Вихоревка", "ranges": [(57, 85)], "track1": "desc"},
    "vih-taishet": {"label": "Вихоревка — Тайшет", "ranges": [(27, 57)], "track1": "desc"},
    "nizh-yurty": {"label": "Нижнеудинск — Юрты", "ranges": [(2, 25)], "track1": "desc"},
    "yurty-nizh": {"label": "Юрты — Нижнеудинск", "ranges": [(2, 25)], "track1": "asc"},
    "bratsk": {"label": "Братский ход", "ranges": [(100, 111)], "track1": "asc"},
    "kor-ust": {"label": "Коршуниха — Усть-Илим", "ranges": [(85, 87), (112, 129)], "track1": "asc"},
    "kor-lena": {"label": "Коршуниха — Лена", "ranges": [(85, 99)], "track1": "asc"},
}

ALL_SPEED_COLUMNS = [
    (4, "Пасс."), (5, "Ускор. груз."), (6, "Груз."), (7, "Порож."),
    (8, "Пасс. ст."), (9, "Ускор. ст."), (10, "Груз. ст."), (11, "Порож. ст."),
    (12, "Пасс. п/о"), (13, "Груз. п/о"),
]


def txt(value):
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def slug_id(value):
    return re.sub(r"[^a-z0-9а-яё]+", "-", value.lower()).strip("-")


def is_speed(value):
    return value is not None and txt(value) != ""


def route_txt(value):
    text = txt(value)
    # 248 км.1 пк 0 м / 248 км 1 пк 25 м -> 248 км 1пк
    return re.sub(r"(\d+)\s*км\s*[.,]?\s*(\d+)\s*пк(?:\s*\d+\s*м)?", r"\1 км \2пк", text, flags=re.IGNORECASE)


def route_limits(value):
    points = re.findall(r"\d+\s*км\s*\d+пк", route_txt(value), flags=re.IGNORECASE)
    if len(points) >= 2:
        return f" ({points[0]} / {points[-1]})"
    return ""


def load_groups():
    wb = load_workbook(SOURCE, data_only=True, read_only=False)
    ws = wb.active
    section_rows = set()
    merged_values = {}
    for rng in ws.merged_cells.ranges:
        if rng.min_row == rng.max_row and rng.min_col == 1 and rng.max_col == 15 and rng.min_row >= 4:
            section_rows.add(rng.min_row)
        if rng.max_col <= 15 and rng.min_row >= 4:
            top_left = ws.cell(rng.min_row, rng.min_col).value
            for merged_row in range(rng.min_row, rng.max_row + 1):
                for merged_col in range(rng.min_col, rng.max_col + 1):
                    merged_values[(merged_row, merged_col)] = top_left

    groups = []
    current = None
    for row_no in range(4, ws.max_row + 1):
        values = [merged_values.get((row_no, c), ws.cell(row_no, c).value) for c in range(1, 16)]
        if row_no in section_rows:
            title = txt(values[0]) or f"Раздел {row_no}"
            current = {"index": len(groups), "title": title, "row": row_no, "rows": []}
            groups.append(current)
        elif current and any(v not in (None, "") for v in values):
            current["rows"].append({"excel_row": row_no, "v": values})
    return groups


def selected_columns(page):
    return page.get("columns", ALL_SPEED_COLUMNS)


def filtered_rows(group, page):
    rows = []
    is_station = group["title"].lower().startswith("станц")
    for row in group["rows"]:
        v = row["v"]
        # Rows with an explicit path belong only to that path. Station rows often
        # have a blank path in the source and therefore remain visible on both paths.
        if page.get("track") and txt(v[1]) not in ("", page["track"]):
            continue
        cols = selected_columns(page)
        speeds = [(label, txt(v[idx])) for idx, label in cols if is_speed(v[idx])]

        # The source sheet has no separate values in the plain-empty station
        # column for many stations. In those rows the freight station value is
        # the applicable value; reception/departure tracks are likewise kept
        # in the common freight station column. Use those station-only
        # fallbacks so the empty pages retain the same station structure and
        # formatting as the other train categories. Pergon values remain taken
        # strictly from their own category columns.
        if is_station and not speeds and page.get("category") == "empty":
            fallback = 9 if txt(v[1]) in ("1", "2") else 13
            if is_speed(v[fallback]):
                speeds = [("Станция", txt(v[fallback]))]
        elif is_station and not speeds and page.get("category") == "18-100":
            if is_speed(v[13]):
                speeds = [("Станция", txt(v[13]))]
        if speeds:
            rows.append((row, speeds))
    return rows


def render_speed_summary(rows, general_only=False):
    ordered = []
    seen = set()
    source_rows = rows[:1] if general_only else rows
    for _, speeds in source_rows:
        for _, speed in speeds:
            if speed not in seen:
                seen.add(speed)
                ordered.append(speed)
    def key(s):
        try:
            return (0, float(s.replace(",", ".")))
        except ValueError:
            return (1, s)
    ordered.sort(key=key)
    return "".join(f'<span class="speed">{escape(s)} <small>км/ч</small></span>' for s in ordered)


def render_station_speed_summary(rows):
    if not rows:
        return ""
    primary = []
    for _, speed in rows[0][1]:
        if speed not in primary:
            primary.append(speed)
    switch_speeds = []
    for _, speeds in rows[1:]:
        for _, speed in speeds:
            if speed not in switch_speeds:
                switch_speeds.append(speed)
    def number(s):
        try:
            return float(s.replace(",", "."))
        except ValueError:
            match = re.search(r"\d+(?:[.,]\d+)?", s)
            return float(match.group(0).replace(",", ".")) if match else float("-inf")
    station_speed = max(primary, key=number) if primary else ""
    max_switch = max(switch_speeds, key=number) if switch_speeds else ""
    has_lower_restrictions = bool(max_switch) and any(number(speed) < number(max_switch) for speed in switch_speeds)
    parts = []
    if station_speed:
        parts.append(f'<span class="speed">{escape(station_speed)} км/ч</span>')
    if max_switch:
        parts.append(f'<span class="switch-speed">стрелки {escape(max_switch)} км/ч</span>')
    if has_lower_restrictions:
        parts.append('<span class="warning" title="Есть дополнительные ограничения" aria-label="Есть ограничения">!</span>')
    return "".join(parts)


def render_group(group, rows, page):
    body = []
    kind = "station" if group["title"].lower().startswith("станц") else "section"
    if kind == "station":
        display_rows = rows[1:] if rows and route_limits(rows[0][0]["v"][2]) else rows
    else:
        general_speeds = {speed for _, speed in rows[0][1]} if rows else set()
        display_rows = [item for item in rows[1:] if any(speed not in general_speeds for _, speed in item[1])]
    for row, speeds in display_rows:
        v = row["v"]
        badges = "".join(f'<span class="value"><b>{escape(speed)}</b></span>' for _, speed in speeds)
        body.append(
            '<tr>'
            f'<td class="name">{escape(route_txt(v[2]))}</td>'
            f'<td class="values">{badges}</td>'
            f'<td class="note">{escape(txt(v[14]))}</td>'
            '</tr>'
        )
    open_attr = "" if kind == "station" else " open"
    limits = route_limits(rows[0][0]["v"][2]) if rows else ""
    display_title = re.sub(r"^Станция\s+", "Ст. ", group["title"], flags=re.IGNORECASE) if kind == "station" else group["title"]
    summary_speeds = render_station_speed_summary(rows) if kind == "station" else render_speed_summary(rows, general_only=True)
    segment_keys = [key for key, section in SECTIONS.items() if any(start <= group["index"] <= end for start, end in section["ranges"])]
    return f'''<details class="group {kind}" data-index="{group['index']}" data-segments="{' '.join(segment_keys)}" id="{slug_id(group['title'])}"{open_attr}>
      <summary><span class="chevron">›</span><span class="group-title">{escape(display_title + limits)}</span>
      <span class="summary-speeds">{summary_speeds}</span><span class="count">{len(display_rows)}</span></summary>
      <div class="table-wrap"><table><thead><tr><th>Участок, станция или путь</th><th>Скорость</th><th>Примечание</th></tr></thead>
      <tbody>{''.join(body)}</tbody></table></div>
    </details>'''


CSS = r'''
:root{--ink:#17202a;--muted:#617080;--line:#d9e0e6;--blue:#125d98;--blue2:#e9f3fb;--nav:#102f49;--speed:#eaf8ee;--speed-ink:#116530;--white:#fff;--head1:64px;--head2:50px;--head3:48px;--head4:44px}
*{box-sizing:border-box}html{scroll-behavior:smooth}body{margin:0;background:#f4f6f8;color:var(--ink);font:15px/1.4 "Segoe UI",Arial,sans-serif}
.sticky{position:sticky;z-index:20}.bar1{top:0;height:var(--head1);background:var(--nav);color:#fff;display:flex;align-items:center;padding:8px 22px;gap:18px}.bar1 h1{margin:0;font-size:22px}.bar1 p{margin:2px 0 0;color:#c8d9e6;font-size:13px}.home{color:white;text-decoration:none;font-size:24px}.bar2{top:var(--head1);height:var(--head2);background:#fff;border-bottom:1px solid var(--line);display:flex;align-items:center;padding:6px 18px;gap:7px;overflow-x:auto}.bar2 a{white-space:nowrap;text-decoration:none;color:#244154;background:#edf2f5;padding:8px 12px;border-radius:8px}.bar2 a.cat-passenger{background:#c83b3b;color:#fff}.bar2 a.cat-freight{background:#198754;color:#fff}.bar2 a.cat-empty{background:#d97819;color:#fff}.bar2 a.cat-18-100{background:#2368b5;color:#fff}.bar2 a.active{outline:3px solid #132c3d;outline-offset:1px;font-weight:750}.bar3{top:calc(var(--head1) + var(--head2));height:var(--head3);background:#f8fafb;border-bottom:1px solid var(--line);display:flex;align-items:center;padding:7px 18px;gap:10px;box-shadow:0 2px 5px #00000012}.bar3 input{flex:1;min-width:180px;max-width:480px;width:auto;height:34px;border:1px solid #bcc9d3;border-radius:8px;padding:0 12px;font:inherit}.bar3 button{height:34px;border:0;border-radius:8px;padding:0 12px;cursor:pointer;background:#dfe8ee;color:#243746;white-space:nowrap}.bar3 .icon-button{width:38px;padding:0;display:grid;place-items:center;font:bold 22px/1 "Segoe UI",Arial}.bar3 .icon-button svg{width:22px;height:22px}.bar3 .icon-button.active{background:#125d98;color:#fff}.status{margin-left:auto;color:var(--muted);white-space:nowrap}.bar4{top:calc(var(--head1) + var(--head2) + var(--head3));height:var(--head4);display:grid;grid-template-columns:minmax(300px,42fr) minmax(260px,28fr) minmax(230px,30fr);align-items:center;background:#d4e4ef;color:#203d52;border-bottom:1px solid #b9cedd;padding:0 26px;font-size:12px;font-weight:750;text-transform:uppercase;letter-spacing:.03em;box-shadow:0 2px 5px #00000015}.bar4 span{padding:0 10px}.bar4 .note-head{visibility:hidden}.notes-visible .bar4 .note-head{visibility:visible}
main{max-width:1500px;margin:10px auto 40px;padding:0 16px}.group{background:#fff;border:1px solid var(--line);border-radius:8px;margin:5px 0;overflow:clip;box-shadow:0 1px 2px #00000008}.group[hidden]{display:none}.group summary{min-height:40px;display:flex;align-items:center;gap:7px;padding:4px 10px;cursor:pointer;list-style:none}.group summary::-webkit-details-marker{display:none}.group[open] summary{background:#eef5fa;border-bottom:1px solid var(--line)}.chevron{font-size:23px;line-height:1;color:var(--blue);transform:rotate(0);transition:.15s}.group[open] .chevron{transform:rotate(90deg)}.group-title{font-weight:700;min-width:240px;flex:1;line-height:1.18}.summary-speeds{display:flex;align-items:center;gap:4px;justify-content:flex-end;flex-wrap:wrap}.speed,.switch-speed{background:var(--speed);color:var(--speed-ink);font-weight:750;padding:3px 6px;border-radius:6px;line-height:1.15}.switch-speed{background:#fff1dc;color:#754317}.speed small{font-weight:500}.warning{display:inline-grid;place-items:center;width:23px;height:21px;padding-top:5px;background:#ffd23f;color:#151515;font-weight:900;font-size:13px;line-height:1;clip-path:polygon(50% 0,100% 100%,0 100%)}.count{min-width:24px;height:24px;border-radius:12px;background:#edf1f4;color:var(--muted);display:grid;place-items:center;font-size:11px}.table-wrap{overflow-x:auto}table{width:100%;border-collapse:collapse;table-layout:fixed;min-width:760px}thead{display:none}th,td{border-bottom:1px solid #e7ecef;text-align:left;padding:3px 8px;vertical-align:middle;line-height:1.18}.name{width:42%}.values{width:28%}.note{width:30%;visibility:hidden;color:#475866}.notes-visible .note{visibility:visible}.value{display:inline-flex;align-items:center;margin:0 3px 0 0;padding:2px 5px;border-radius:5px;background:var(--blue2);color:#174b70;line-height:1.1}.value b{font-size:15px}.empty-state{text-align:center;color:var(--muted);padding:80px 20px}.intro{background:#fff;border:1px solid var(--line);border-radius:8px;padding:8px 12px;margin-bottom:8px;color:#435463}
.group.station{border-color:#a95c29}.group.station>summary,.group.station[open]>summary{background:#c77a3b;color:#151515;border-bottom-color:#9d4f20}.group.station .chevron{color:#151515}.group.station .count{background:#f4dfcf;color:#151515}
@media(max-width:760px){:root{--head1:72px;--head2:48px;--head3:46px;--head4:42px}.bar1{padding:8px 12px}.bar1 h1{font-size:18px}.bar1 p{display:none}.bar2,.bar3{padding-left:10px}.bar3 button:not(#toggleNotes){display:none}.bar4{padding:0 10px;grid-template-columns:300px 250px 230px;overflow:hidden}.status{font-size:12px}.group-title{min-width:0}.group summary{align-items:flex-start}.summary-speeds{max-width:52%}.speed{font-size:13px}.count{display:none}}
@media print{.bar2,.bar3{display:none}.bar1,.bar4{position:static}.group{break-inside:avoid}.group:not([open]) .table-wrap{display:block}main{max-width:none}}
'''

JS = r'''
const groups=[...document.querySelectorAll('.group')];
const q=document.getElementById('q'), status=document.getElementById('status');
const params=new URLSearchParams(location.search), section=params.get('section')||'';
const sectionNames=SECTION_NAMES;
const track=PAGE_TRACK, track1Directions=ORDER_DIRECTIONS;
if(section&&sectionNames[section]){document.getElementById('chosenSection').textContent=sectionNames[section];document.querySelectorAll('nav a[data-category]').forEach(a=>a.href+=`?section=${encodeURIComponent(section)}`);const base=track1Directions[section]==='desc'?-1:1, direction=track==='1'?base:-base;const selectedGroups=groups.filter(g=>g.dataset.segments.split(' ').includes(section));selectedGroups.sort((a,b)=>direction*(Number(a.dataset.index)-Number(b.dataset.index))).forEach(g=>g.parentElement.appendChild(g));if(direction===-1){selectedGroups.forEach(g=>{const el=g.querySelector('.group-title'),text=el.textContent;if(g.classList.contains('section')){const m=text.match(/^(Перегон\s+)(.+?)\s+-\s+(.+?)(?:\s+\(([^/]+)\s*\/\s*([^)]+)\))?$/);if(m)el.textContent=`${m[1]}${m[3]} - ${m[2]}${m[4]&&m[5]?` (${m[5].trim()} / ${m[4].trim()})`:''}`;const tbody=g.querySelector('tbody');if(tbody){Array.from(tbody.children).reverse().forEach(row=>{const cell=row.querySelector('.name');if(cell){const range=cell.textContent.match(/^(\d+\s*км\s*\d+пк)\s*-\s*(\d+\s*км\s*\d+пк)$/i);if(range)cell.textContent=`${range[2]} - ${range[1]}`}tbody.appendChild(row)})}}else if(g.classList.contains('station')){const m=text.match(/^(Ст\.\s+.+?)\s+\(([^/]+)\s*\/\s*([^)]+)\)$/);if(m)el.textContent=`${m[1]} (${m[3].trim()} / ${m[2].trim()})`}})}}
function apply(){const s=q.value.trim().toLocaleLowerCase('ru');let n=0;groups.forEach(g=>{const inSection=!section||g.dataset.segments.split(' ').includes(section);const ok=inSection&&(!s||g.textContent.toLocaleLowerCase('ru').includes(s));g.hidden=!ok;if(ok)n++});status.textContent=`Разделов: ${n}`}
q.addEventListener('input',apply);document.getElementById('openAll').onclick=()=>groups.filter(g=>!g.hidden).forEach(g=>g.open=true);document.getElementById('closeAll').onclick=()=>groups.forEach(g=>g.open=false);document.getElementById('toggleNotes').onclick=e=>{const visible=document.body.classList.toggle('notes-visible');e.currentTarget.classList.toggle('active',visible);e.currentTarget.title=visible?'Скрыть примечания':'Показать примечания';e.currentTarget.setAttribute('aria-label',e.currentTarget.title)};apply();
'''

PWA_REGISTER = r'''<script>
if ('serviceWorker' in navigator) window.addEventListener('load',async()=>{
  let refreshing=false;
  navigator.serviceWorker.addEventListener('controllerchange',()=>{if(!refreshing){refreshing=true;location.reload()}});
  const registration=await navigator.serviceWorker.register('sw.js',{updateViaCache:'none'});
  if(registration.waiting) registration.waiting.postMessage('SKIP_WAITING');
  registration.addEventListener('updatefound',()=>{const worker=registration.installing;if(worker)worker.addEventListener('statechange',()=>{if(worker.state==='installed'&&navigator.serviceWorker.controller)worker.postMessage('SKIP_WAITING')})});
  registration.update();
});
</script>'''


def nav(active, track):
    entries = [(f"track-{track}-{key}.html", category["title"]) for key, category in CATEGORIES.items()]
    return "".join(f'<a data-category="1" href="{file}" class="cat-{key} {"active" if file == active else ""}">{label}</a>' for key, (file, label) in zip(CATEGORIES, entries))


def page_html(filename, page, groups):
    rendered = []
    source_rows = 0
    for group in groups:
        rows = filtered_rows(group, page)
        if rows:
            rendered.append(render_group(group, rows, page))
            source_rows += len(rows)
    content = "".join(rendered) or '<div class="empty-state">По этому фильтру данных нет.</div>'
    section_names = json.dumps({key: value["label"] for key, value in SECTIONS.items()}, ensure_ascii=False)
    directions = json.dumps({key: value["track1"] for key, value in SECTIONS.items()}, ensure_ascii=False)
    script = JS.replace("SECTION_NAMES", section_names).replace("ORDER_DIRECTIONS", directions).replace("PAGE_TRACK", json.dumps(page["track"]))
    return f'''<!doctype html><html lang="ru"><head><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1"><meta name="theme-color" content="#123a59"><link rel="manifest" href="manifest.webmanifest"><link rel="icon" href="icon.svg" type="image/svg+xml"><title>{escape(page['title'])}</title><style>{CSS}</style></head><body>
<header class="bar1 sticky"><a class="home" href="index.html" title="Все фильтры">⌂</a><div><h1>{escape(page['title'])}</h1><p><span id="chosenSection">Все участки</span> · {escape(page['subtitle'])}</p></div></header>
<nav class="bar2 sticky"><a href="index.html">← Выбор пути</a>{nav(filename, page['track'])}</nav>
<div class="bar3 sticky"><input id="q" type="search" placeholder="Поиск по станции, перегону, примечанию…"><button class="icon-button" id="openAll" title="Развернуть всё" aria-label="Развернуть всё">=</button><button class="icon-button" id="closeAll" title="Свернуть всё" aria-label="Свернуть всё">−</button><button class="icon-button" id="toggleNotes" title="Показать примечания" aria-label="Показать примечания"><svg viewBox="0 0 24 24" aria-hidden="true"><path d="M5 3.5h10l4 4V20.5H5z" fill="none" stroke="currentColor" stroke-width="1.7"/><path d="M15 3.5v4h4M8 12h5M8 15h4" fill="none" stroke="currentColor" stroke-width="1.5"/><path d="M13.5 17.7c2.1-4 4.2-6 7-7.2-1 2.8-2.8 5.4-6.7 7.7l-1.8 2.1" fill="none" stroke="currentColor" stroke-width="1.5" stroke-linecap="round"/></svg></button><span class="status" id="status"></span></div>
<div class="bar4 sticky"><span>Участок, станция или путь</span><span>Скорость</span><span class="note-head">Примечание</span></div>
<main><div class="intro">Станции изначально свернуты, перегоны раскрыты. Скорости видны в заголовке; нажмите на станцию, чтобы открыть строки.</div>{content}</main><script>{script}</script>{PWA_REGISTER}</body></html>'''


def index_html():
    section_buttons = "".join(f'<button type="button" data-group="section" data-value="{key}">{escape(value["label"])}</button>' for key, value in SECTIONS.items())
    train_buttons = "".join(f'<button class="cat-{key}" type="button" data-group="train" data-value="{key}">{escape(value["title"])}</button>' for key, value in CATEGORIES.items())
    return f'''<!doctype html><html lang="ru"><head><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1"><meta name="theme-color" content="#123a59"><link rel="manifest" href="manifest.webmanifest"><link rel="icon" href="icon.svg" type="image/svg+xml"><title>Выбор скоростей</title><style>{CSS}.chooser{{max-width:1050px;margin:0 auto;padding:38px 18px 60px}}.chooser-head{{display:flex;align-items:flex-start;justify-content:space-between;gap:18px}}.chooser-actions{{display:flex;gap:9px}}.chooser h1{{font-size:36px;margin:0 0 8px}}.lead{{color:var(--muted);margin:0 0 28px}}.head-button{{display:grid;place-items:center;flex:0 0 48px;width:48px;height:48px;border:1px solid #b9c9d4;border-radius:11px;background:#fff;color:#123a59;text-decoration:none;font:bold 14px "Segoe UI",Arial;box-shadow:0 2px 8px #00000012;cursor:pointer}}.head-button:hover{{background:#eaf3f8;border-color:#6f9fbe}}.head-button.install-ready{{background:#198754;color:#fff;border-color:#198754}}.step{{background:#fff;border:1px solid var(--line);border-radius:14px;padding:20px;margin:14px 0}}.step h2{{margin:0 0 14px;font-size:20px}}.options{{display:flex;flex-wrap:wrap;gap:10px}}.options button{{border:2px solid #cad5dc;background:#f2f5f7;color:#263b49;border-radius:10px;padding:11px 15px;font:600 15px/1.2 "Segoe UI",Arial;cursor:pointer;transition:.12s}}.options button:hover{{filter:brightness(.93)}}.options button.cat-passenger{{background:#c83b3b;border-color:#c83b3b;color:#fff}}.options button.cat-freight{{background:#198754;border-color:#198754;color:#fff}}.options button.cat-empty{{background:#d97819;border-color:#d97819;color:#fff}}.options button.cat-18-100{{background:#2368b5;border-color:#2368b5;color:#fff}}.options button.selected{{outline:3px solid #132c3d;outline-offset:2px;box-shadow:0 3px 10px #0000002f}}.go{{display:block;width:100%;margin-top:20px;padding:15px;border:0;border-radius:11px;background:#125d98;color:#fff;font:bold 19px "Segoe UI",Arial;cursor:pointer}}.go:disabled{{background:#aeb9c1;cursor:not-allowed}}.hint{{color:#6a7883;font-size:13px;margin:10px 0 0}}</style></head><body><main class="chooser"><div class="chooser-head"><div><h1>Выбор скоростей</h1><p class="lead">Выберите участок, путь и тип поезда.</p></div><div class="chooser-actions"><a class="head-button" href="source.xlsx" title="Открыть исходный Excel" aria-label="Открыть исходный Excel">🖨</a><button class="head-button" id="installPwa" title="Установить для офлайн-доступа" aria-label="Установить приложение">PWA</button></div></div><section class="step"><h2>1. Участок</h2><div class="options">{section_buttons}</div></section><section class="step"><h2>2. Номер пути</h2><div class="options"><button type="button" data-group="track" data-value="1">1 путь</button><button type="button" data-group="track" data-value="2">2 путь</button></div></section><section class="step"><h2>3. Тип поезда</h2><div class="options">{train_buttons}</div></section><button class="go" id="go" disabled>Поехали</button><p class="hint">Выберите один вариант в каждом разделе.</p></main><script>const selected={{section:'',track:'',train:''}},go=document.getElementById('go');document.querySelectorAll('[data-group]').forEach(btn=>btn.onclick=()=>{{const group=btn.dataset.group;document.querySelectorAll(`[data-group="${{group}}"]`).forEach(b=>b.classList.remove('selected'));btn.classList.add('selected');selected[group]=btn.dataset.value;go.disabled=!(selected.section&&selected.track&&selected.train)}});go.onclick=()=>{{location.href=`track-${{selected.track}}-${{selected.train}}.html?section=${{encodeURIComponent(selected.section)}}`}};let installPrompt=null;const installButton=document.getElementById('installPwa');window.addEventListener('beforeinstallprompt',e=>{{e.preventDefault();installPrompt=e;installButton.classList.add('install-ready')}});installButton.onclick=async()=>{{if(installPrompt){{installPrompt.prompt();await installPrompt.userChoice;installPrompt=null;installButton.classList.remove('install-ready')}}else if(/iphone|ipad|ipod/i.test(navigator.userAgent))alert('На iPhone: нажмите «Поделиться», затем «На экран Домой».');else alert('Откройте сайт по защищённой ссылке и выберите установку приложения в меню браузера.')}};window.addEventListener('appinstalled',()=>installButton.hidden=true);</script>{PWA_REGISTER}</body></html>'''


def track_landing(track):
    cards = []
    for key, category in CATEGORIES.items():
        cards.append(f'<a class="card" href="track-{track}-{key}.html"><h2>{escape(category["title"])}</h2><p>{escape(category["subtitle"])}</p><span>Открыть →</span></a>')
    return f'''<!doctype html><html lang="ru"><head><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1"><title>{track} путь</title><style>{CSS}.landing{{max-width:1050px;margin:0 auto;padding:34px 18px}}.back{{text-decoration:none;color:var(--blue);font-weight:700}}.landing h1{{font-size:34px;margin:18px 0 6px}}.landing>p{{color:var(--muted);margin:0 0 26px}}.cards{{display:grid;grid-template-columns:repeat(2,minmax(270px,1fr));gap:14px}}.card{{display:block;background:#fff;border:1px solid var(--line);border-radius:12px;padding:20px;text-decoration:none;color:var(--ink);box-shadow:0 2px 8px #0000000b;transition:.15s}}.card:hover{{transform:translateY(-2px);border-color:#8bb6d4}}.card h2{{margin:0 0 8px;color:var(--blue)}}.card p{{min-height:44px;color:var(--muted)}}.card span{{font-weight:700}}@media(max-width:650px){{.cards{{grid-template-columns:1fr}}}}</style></head><body><div class="landing"><a class="back" href="index.html">← Изменить путь</a><h1>{track} путь</h1><p>Выберите вид движения.</p><div class="cards">{''.join(cards)}</div></div></body></html>'''


def main():
    groups = load_groups()
    OUT.mkdir(parents=True, exist_ok=True)
    report = {"source": str(SOURCE), "groups_total": len(groups), "pages": {}}
    for filename, page in PAGES.items():
        html = page_html(filename, page, groups)
        (OUT / filename).write_text(html, encoding="utf-8")
        report["pages"][filename] = {"bytes": len(html.encode("utf-8")), "details": html.count('<details class="group')}
    (OUT / "index.html").write_text(index_html(), encoding="utf-8")
    manifest = {
        "name": "Скорости по станциям",
        "short_name": "Скорости",
        "description": "Скорости движения по станциям и перегонам",
        "start_url": "./index.html",
        "scope": "./",
        "display": "standalone",
        "background_color": "#eef3f6",
        "theme_color": "#123a59",
        "lang": "ru",
        "icons": [{"src": "icon.svg", "sizes": "any", "type": "image/svg+xml", "purpose": "any maskable"}],
    }
    (OUT / "manifest.webmanifest").write_text(json.dumps(manifest, ensure_ascii=False, indent=2), encoding="utf-8")
    icon = '''<svg xmlns="http://www.w3.org/2000/svg" viewBox="0 0 512 512"><rect width="512" height="512" rx="92" fill="#123a59"/><path d="M95 350h322v54H95zM128 108h256v188H128z" fill="#fff"/><path d="M160 147h192v28H160zm0 59h192v28H160zm0 59h125v28H160z" fill="#198754"/><circle cx="372" cy="350" r="54" fill="#d97819"/><path d="m348 350 16 17 34-39" fill="none" stroke="#fff" stroke-width="18" stroke-linecap="round" stroke-linejoin="round"/></svg>'''
    (OUT / "icon.svg").write_text(icon, encoding="utf-8")
    cache_files = ["./", "index.html", "manifest.webmanifest", "icon.svg", "source.xlsx", *PAGES.keys()]
    sw = f'''const CACHE="speeds-v3";const FILES={json.dumps(cache_files, ensure_ascii=False)};self.addEventListener("install",e=>e.waitUntil(caches.open(CACHE).then(c=>c.addAll(FILES)).then(()=>self.skipWaiting())));self.addEventListener("activate",e=>e.waitUntil(caches.keys().then(keys=>Promise.all(keys.filter(k=>k!==CACHE).map(k=>caches.delete(k)))).then(()=>self.clients.claim())));self.addEventListener("message",e=>{{if(e.data==="SKIP_WAITING")self.skipWaiting()}});self.addEventListener("fetch",e=>{{if(e.request.method!=="GET")return;if(e.request.mode==="navigate"){{e.respondWith(fetch(e.request).then(r=>{{const copy=r.clone();caches.open(CACHE).then(c=>c.put(e.request,copy));return r}}).catch(()=>caches.match(e.request,{{ignoreSearch:true}}).then(hit=>hit||caches.match("index.html"))));return}}e.respondWith(caches.match(e.request).then(hit=>hit||fetch(e.request).then(r=>{{const copy=r.clone();caches.open(CACHE).then(c=>c.put(e.request,copy));return r}})))}});'''
    (OUT / "sw.js").write_text(sw, encoding="utf-8")
    shutil.copy2(SOURCE, OUT / "source.xlsx")
    (OUT / "validation.json").write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps(report, ensure_ascii=True))


if __name__ == "__main__":
    main()
